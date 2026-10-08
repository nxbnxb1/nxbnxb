"""Generate demo documents (Vietnamese + English) for the Extract workflow:

    samples/bao_cao_digital.pdf   PDF with a text layer
    samples/bao_cao_scan.pdf      the same page as a skewed, noisy scan (image only)
    samples/bao_cao.docx          Word version with headings, list and a merged-cell table
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import numpy as np
import pymupdf
from docx import Document
from PIL import Image, ImageFilter

FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
]
TITLE = "BÁO CÁO KẾT QUẢ KINH DOANH NĂM 2025"
PARAGRAPHS = [
    "Năm 2025, doanh thu thuần của Công ty đạt 1.234,5 tỷ đồng, tăng 12,3% so với năm 2024. "
    "Lợi nhuận sau thuế đạt 156,8 tỷ đồng, vượt 4,2% kế hoạch được Đại hội đồng cổ đông thông qua.",
    "The Board of Directors proposes a cash dividend of 15% for the 2025 financial year.",
]
BULLETS = ["Mở rộng hệ thống phân phối tại 12 tỉnh, thành phố.", "Đầu tư dây chuyền sản xuất mới trị giá 85 tỷ đồng."]
TABLE = [["Chỉ tiêu", "Năm 2024", "Năm 2025", "Tăng trưởng"],
         ["Doanh thu thuần", "1.099,2", "1.234,5", "12,3%"],
         ["Lợi nhuận gộp", "312,4", "355,0", "13,6%"],
         ["Lợi nhuận sau thuế", "140,1", "156,8", "11,9%"]]


def font_path() -> str:
    for path in FONT_CANDIDATES:
        if Path(path).exists():
            return path
    raise SystemExit("no Vietnamese-capable font found (install fonts-dejavu-core)")


def digital_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_font(fontname="vn", fontfile=font_path())
    page.insert_text((60, 70), TITLE, fontsize=17, fontname="vn")
    page.insert_text((60, 110), "1. Kết quả hoạt động", fontsize=13, fontname="vn")
    box = pymupdf.Rect(60, 122, 535, 230)
    page.insert_textbox(box, "\n".join(PARAGRAPHS), fontsize=10.5, fontname="vn")
    y = 245
    for row_index, row in enumerate(TABLE):
        for col, value in enumerate(row):
            cell = pymupdf.Rect(60 + col * 118, y, 60 + (col + 1) * 118, y + 22)
            page.draw_rect(cell, width=0.6)
            page.insert_text((cell.x0 + 5, cell.y0 + 15), value, fontsize=9.5, fontname="vn")
        y += 22
    page.insert_text((60, y + 35), "2. Kế hoạch năm 2026", fontsize=13, fontname="vn")
    for i, bullet in enumerate(BULLETS):
        page.insert_text((70, y + 58 + i * 16), f"- {bullet}", fontsize=10.5, fontname="vn")
    page.insert_text((290, 815), "1", fontsize=9, fontname="vn")
    return doc.tobytes(garbage=3, deflate=True)


def scanned_pdf(digital: bytes) -> bytes:
    src = pymupdf.open(stream=digital, filetype="pdf")
    pix = src[0].get_pixmap(dpi=200, alpha=False)
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    image = image.rotate(1.2, resample=Image.Resampling.BICUBIC, fillcolor="white").filter(ImageFilter.GaussianBlur(0.6))
    arr = np.asarray(image, dtype=np.float32) * 0.92 + 12
    arr += np.random.default_rng(0).normal(0, 6, arr.shape)
    image = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))
    buf = io.BytesIO()
    image.save(buf, format="JPEG", quality=70)
    out = pymupdf.open()
    page = out.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=buf.getvalue())
    return out.tobytes(garbage=3, deflate=True)


def docx() -> bytes:
    doc = Document()
    doc.add_heading(TITLE, level=0)
    doc.add_heading("1. Kết quả hoạt động", level=1)
    for paragraph in PARAGRAPHS:
        doc.add_paragraph(paragraph)
    table = doc.add_table(rows=len(TABLE) + 1, cols=4)
    table.style = "Table Grid"
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Chỉ tiêu (tỷ đồng)"
    table.cell(0, 1).merge(table.cell(0, 2)).text = "Giá trị"
    table.cell(0, 3).merge(table.cell(1, 3)).text = "Tăng trưởng"
    table.cell(1, 1).text, table.cell(1, 2).text = "Năm 2024", "Năm 2025"
    for r, row in enumerate(TABLE[1:], start=2):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.add_heading("2. Kế hoạch năm 2026", level=1)
    for bullet in BULLETS:
        doc.add_paragraph(bullet, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def main() -> None:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "samples")
    out.mkdir(parents=True, exist_ok=True)
    digital = digital_pdf()
    (out / "bao_cao_digital.pdf").write_bytes(digital)
    (out / "bao_cao_scan.pdf").write_bytes(scanned_pdf(digital))
    (out / "bao_cao.docx").write_bytes(docx())
    print("\n".join(str(p) for p in sorted(out.iterdir())))


if __name__ == "__main__":
    main()
