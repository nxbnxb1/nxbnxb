"""Word (.docx) reading with python-docx: structure is native, so no OCR is needed."""

from __future__ import annotations

import io
import logging
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from PIL import Image

from ..models import RegionType
from ..tables import Cell, Table
from ..textutil import normalize

log = logging.getLogger(__name__)

_HEADING_RE = re.compile(r"(heading|tiêu đề|título|titre|überschrift|überschrift)\s*(\d)", re.IGNORECASE)


@dataclass
class DocxBlock:
    type: RegionType
    content: str = ""
    level: int | None = None
    table: Table | None = None
    image: Image.Image | None = None
    locator: str = ""


def convert_doc_to_docx(data: bytes) -> bytes:
    """Legacy .doc → .docx through LibreOffice, when it is installed."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ValueError("Legacy .doc files need LibreOffice (soffice) for conversion; please save as .docx")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "input.doc"
        src.write_bytes(data)
        subprocess.run(
            [soffice, "--headless", "--convert-to", "docx", "--outdir", tmp, str(src)],
            check=True,
            capture_output=True,
            timeout=300,
        )
        return (Path(tmp) / "input.docx").read_bytes()


class DocxReader:
    def __init__(self, data: bytes) -> None:
        self.doc = Document(io.BytesIO(data))
        self._num_formats = self._numbering_formats()

    def blocks(self) -> list[DocxBlock]:
        out: list[DocxBlock] = []
        p_index = t_index = 0
        for item in self.doc.iter_inner_content():
            if isinstance(item, Paragraph):
                p_index += 1
                out.extend(self._paragraph(item, f"body/p[{p_index}]"))
            elif isinstance(item, DocxTable):
                t_index += 1
                table = self._table(item)
                if table.n_cells:
                    out.append(DocxBlock(RegionType.TABLE, table=table, locator=f"body/tbl[{t_index}]"))
        return _merge_lists(out)

    # --- paragraphs ---------------------------------------------------------------

    def _paragraph(self, p: Paragraph, locator: str) -> list[DocxBlock]:
        blocks: list[DocxBlock] = []
        for i, image in enumerate(self._images(p)):
            blocks.append(DocxBlock(RegionType.IMAGE, content=image[1], image=image[0], locator=f"{locator}/img[{i + 1}]"))
        math = p._p.xpath('.//*[local-name()="oMath"]')
        text = normalize(p.text)
        if math and not text:
            formula = normalize(" ".join("".join(m.xpath('.//*[local-name()="t"]/text()')) for m in math))
            if formula:
                blocks.insert(0, DocxBlock(RegionType.FORMULA, content=formula, locator=locator))
            return blocks
        if not text:
            return blocks
        style = (p.style.name if p.style is not None else "") or ""
        level = self._outline_level(p, style)
        if style.lower() == "title":
            block = DocxBlock(RegionType.TITLE, content=text, level=1, locator=locator)
        elif style.lower() == "subtitle":
            block = DocxBlock(RegionType.HEADING, content=text, level=2, locator=locator)
        elif level is not None:
            block = DocxBlock(RegionType.HEADING, content=text, level=min(6, level + 2), locator=locator)
        elif style.lower() == "caption":
            block = DocxBlock(RegionType.CAPTION, content=text, locator=locator)
        else:
            marker = self._list_marker(p, style)
            if marker is not None:
                block = DocxBlock(RegionType.LIST, content=marker + text, locator=locator)
            else:
                block = DocxBlock(RegionType.TEXT, content=text, locator=locator)
        blocks.insert(0, block)
        return blocks

    def _outline_level(self, p: Paragraph, style: str) -> int | None:
        match = _HEADING_RE.search(style)
        if match:
            return int(match.group(2)) - 1
        values = p._p.xpath("./w:pPr/w:outlineLvl/@w:val")
        style_el = p.style.element if p.style is not None else None
        while not values and style_el is not None:
            values = style_el.xpath("./w:pPr/w:outlineLvl/@w:val")
            based = style_el.xpath("./w:basedOn/@w:val")
            style_el = self._style_by_id(based[0]) if based else None
        if values and values[0].isdigit() and int(values[0]) < 9:
            return int(values[0])
        return None

    def _style_by_id(self, style_id: str):
        for style in self.doc.styles.element.xpath("./w:style"):
            if style.get(qn("w:styleId")) == style_id:
                return style
        return None

    def _numbering_formats(self) -> dict[tuple[str, str], str]:
        """(numId, ilvl) → numFmt ('bullet', 'decimal', ...)."""
        try:
            numbering = self.doc.part.numbering_part.element
        except (KeyError, NotImplementedError, AttributeError):
            return {}
        abstract: dict[str, dict[str, str]] = {}
        for an in numbering.xpath("./w:abstractNum"):
            levels = {lvl.get(qn("w:ilvl")): (lvl.xpath("./w:numFmt/@w:val") or ["bullet"])[0] for lvl in an.xpath("./w:lvl")}
            abstract[an.get(qn("w:abstractNumId"))] = levels
        out = {}
        for num in numbering.xpath("./w:num"):
            ref = num.xpath("./w:abstractNumId/@w:val")
            for ilvl, fmt in abstract.get(ref[0] if ref else "", {}).items():
                out[(num.get(qn("w:numId")), ilvl)] = fmt
        return out

    def _list_marker(self, p: Paragraph, style: str) -> str | None:
        num_id = p._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")
        ilvl = p._p.xpath("./w:pPr/w:numPr/w:ilvl/@w:val") or ["0"]
        if not num_id and not style.lower().startswith("list"):
            return None
        if num_id and num_id[0] == "0":
            return None  # numbering explicitly removed
        depth = int(ilvl[0]) if ilvl[0].isdigit() else 0
        fmt = self._num_formats.get((num_id[0], ilvl[0]), "bullet") if num_id else "bullet"
        if "number" in style.lower():
            fmt = "decimal"
        return "  " * depth + ("- " if fmt in ("bullet", "none") else "1. ")

    def _images(self, p: Paragraph) -> list[tuple[Image.Image, str]]:
        out = []
        for drawing in p._p.xpath('.//*[local-name()="drawing" or local-name()="pict"]'):
            alt = " ".join(v for v in drawing.xpath('.//*[local-name()="docPr"]/@descr') if v)
            for rid in drawing.xpath('.//*[local-name()="blip"]/@*[local-name()="embed"]'):
                part = self.doc.part.related_parts.get(rid)
                if part is None:
                    continue
                try:
                    image = Image.open(io.BytesIO(part.blob))
                    image.load()
                    out.append((image.convert("RGB"), normalize(alt)))
                except Exception as exc:  # EMF/WMF and friends cannot be opened by Pillow
                    log.info("skipping unreadable image %s: %s", rid, exc)
        return out

    # --- tables -------------------------------------------------------------------

    def _table(self, table: DocxTable) -> Table:
        rows: list[list[Cell]] = []
        open_merges: dict[int, Cell] = {}
        for r, tr in enumerate(table._tbl.tr_lst):
            col = tr.grid_before
            row: list[Cell] = []
            is_header = bool(tr.xpath("./w:trPr/w:tblHeader")) or r == 0
            for tc in tr.tc_lst:
                span = tc.grid_span
                if tc.vMerge == "continue":
                    origin = open_merges.get(col)
                    if origin is not None:
                        origin.rowspan += 1
                        col += span
                        continue
                cell = Cell(text=self._cell_text(_Cell(tc, table)), colspan=span, header=is_header)
                if tc.vMerge == "restart":
                    open_merges[col] = cell
                else:
                    open_merges.pop(col, None)
                row.append(cell)
                col += span
            rows.append(row)
        return Table(rows)

    def _cell_text(self, cell: _Cell) -> str:
        parts = [normalize(p.text) for p in cell.paragraphs]
        for nested in cell.tables:
            for row in nested.rows:
                parts.append(" | ".join(normalize(c.text) for c in row.cells))
        return "\n".join(p for p in parts if p)


def _merge_lists(blocks: list[DocxBlock]) -> list[DocxBlock]:
    """Consecutive list paragraphs become one list region."""
    out: list[DocxBlock] = []
    for block in blocks:
        if block.type == RegionType.LIST and out and out[-1].type == RegionType.LIST:
            out[-1].content += "\n" + block.content
            continue
        out.append(block)
    return out
