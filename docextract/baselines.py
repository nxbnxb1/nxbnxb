"""Document conversion systems compared by the benchmark (``docextract bench --system``).

    docextract   this pipeline (layout → router → OCR / tables / formulas / VLM → Markdown); the
                 recognition model follows the settings (product baseline, a fine-tuned model via
                 DOCEXTRACT_OCR_REC_MODEL_DIR, or another stock model via DOCEXTRACT_OCR_REC_MODEL)
    ppstructure  PP-StructureV3, PaddleOCR's own document parser (Markdown output), with the
                 product's PaddleOCR language
    tesseract    Tesseract 5 OCR of every page (tessdata_best, the product's languages), plain text
    text_layer   text already in the file, no OCR: PDF text layer (PyMuPDF), DOCX paragraphs and
                 tables (python-docx); a scan or a picture gives nothing

Every system gets the same file. Formats a system cannot read directly (DOCX for the OCR tools)
are converted to PDF with LibreOffice first, outside the timed part.
"""

from __future__ import annotations

import os
import subprocess
import tempfile
import time
from dataclasses import dataclass, field
from pathlib import Path

from .config import Settings
from .models import DocumentStats

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}


@dataclass
class Converted:
    markdown: str
    stats: DocumentStats = field(default_factory=DocumentStats)


def _timed_stats(pages: int, start: float) -> DocumentStats:
    ms = (time.perf_counter() - start) * 1000
    return DocumentStats(pages_processed=pages, duration_ms=round(ms, 1), ms_per_page=round(ms / max(1, pages), 1))


def _as_pdf(path: Path) -> bytes | None:
    """PDF bytes of a document the OCR baselines must render (None for pictures)."""
    if path.suffix.lower() in IMAGE_SUFFIXES:
        return None
    data = path.read_bytes()
    if path.suffix.lower() == ".pdf":
        return data
    from .preprocessing.office import convert

    return convert(data, path.name, "pdf")


class DocextractSystem:
    name = "docextract"

    def __init__(self, settings: Settings, pipeline=None, use_vlm: bool = True) -> None:
        from .pipeline import DocumentPipeline

        self.pipeline = pipeline or DocumentPipeline(settings)
        self.use_vlm = use_vlm

    def warm_up(self) -> None:
        self.pipeline.warm_up()

    def convert(self, path: Path) -> Converted:
        from .pipeline import ExtractOptions

        result = self.pipeline.process_file(path, ExtractOptions(use_vlm=self.use_vlm))
        return Converted(result.markdown or "", result.stats)


class TextLayerSystem:
    name = "text_layer"

    def warm_up(self) -> None:
        pass

    def convert(self, path: Path) -> Converted:
        start = time.perf_counter()
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            import pymupdf

            with pymupdf.open(path) as doc:
                text = "\n\n".join(page.get_text("text", sort=True) for page in doc)
                pages = doc.page_count
        elif suffix == ".docx":
            import docx

            document = docx.Document(str(path))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                parts += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
            text, pages = "\n\n".join(parts), 1
        else:
            text, pages = "", 1
        return Converted(text, _timed_stats(pages, start))


class TesseractSystem:
    name = "tesseract"

    def __init__(self, langs: str, dpi: int = 300) -> None:
        self.langs = langs
        self.dpi = dpi

    def warm_up(self) -> None:
        subprocess.run(["tesseract", "--version"], capture_output=True, check=True)

    def _ocr(self, png: bytes) -> str:
        cmd = ["tesseract", "stdin", "stdout", "-l", self.langs, "--psm", "3"]
        if os.environ.get("TESSDATA_PREFIX"):
            cmd += ["--tessdata-dir", os.environ["TESSDATA_PREFIX"]]
        return subprocess.run(cmd, input=png, capture_output=True, check=True).stdout.decode("utf-8", "replace")

    def convert(self, path: Path) -> Converted:
        pdf = _as_pdf(path)
        start = time.perf_counter()
        if pdf is None:
            texts = [self._ocr(path.read_bytes())]
        else:
            import pymupdf

            with pymupdf.open(stream=pdf, filetype="pdf") as doc:
                texts = [self._ocr(page.get_pixmap(dpi=self.dpi).tobytes("png")) for page in doc]
        return Converted("\n\n".join(t.strip() for t in texts), _timed_stats(len(texts), start))


class PPStructureSystem:
    name = "ppstructure"

    def __init__(self, lang: str) -> None:
        self.lang = lang
        self._pipeline = None

    def warm_up(self) -> None:
        if self._pipeline is None:
            from paddleocr import PPStructureV3

            self._pipeline = PPStructureV3(
                lang=self.lang,
                use_doc_unwarping=False,
                use_chart_recognition=False,  # needs a VLM; docextract is also measured without one
                enable_mkldnn=False,
            )

    def convert(self, path: Path) -> Converted:
        self.warm_up()
        pdf = _as_pdf(path)
        with tempfile.TemporaryDirectory() as tmp:
            source = path
            if pdf is not None and path.suffix.lower() != ".pdf":
                source = Path(tmp) / f"{path.stem}.pdf"
                source.write_bytes(pdf)
            start = time.perf_counter()
            pages = [res.markdown for res in self._pipeline.predict(input=str(source))]
            joined = self._pipeline.concatenate_markdown_pages(pages)
            if isinstance(joined, tuple):  # some versions also return the images
                joined = joined[0]
            if isinstance(joined, dict):
                joined = joined.get("markdown_texts", "")
            return Converted(str(joined), _timed_stats(len(pages), start))


SYSTEMS = ("docextract", "ppstructure", "tesseract", "text_layer")


# --- supervised worker ---------------------------------------------------------------------


def _memory_total() -> int:
    try:
        for line in Path("/proc/meminfo").read_text().splitlines():
            if line.startswith("MemTotal:"):
                return int(line.split()[1]) * 1024
    except OSError:
        pass
    return 0


def _rss(pid: int) -> int:
    """Resident memory of a process and all its descendants (Linux /proc; 0 elsewhere)."""
    total, todo, seen = 0, [pid], set()
    while todo:
        p = todo.pop()
        if p in seen:
            continue
        seen.add(p)
        try:
            for line in Path(f"/proc/{p}/status").read_text().splitlines():
                if line.startswith("VmRSS:"):
                    total += int(line.split()[1]) * 1024
            for task in Path(f"/proc/{p}/task").iterdir():
                todo += [int(c) for c in (task / "children").read_text().split()]
        except (OSError, ValueError):
            continue
    return total


def _worker(name: str, product: str | None, use_vlm: bool, conn) -> None:
    from .config import Settings

    settings = Settings.from_env(product=product) if product else Settings.from_env()
    system = make_system(name, settings, use_vlm=use_vlm)
    system.warm_up()
    conn.send(("ready", None))
    while (path := conn.recv()) is not None:
        try:
            converted = system.convert(Path(path))
            conn.send(("ok", (converted.markdown, converted.stats.model_dump())))
        except Exception as exc:
            conn.send(("error", f"{type(exc).__name__}: {exc}"[:300]))


class IsolatedSystem:
    """A system run in a supervised worker process. A document that makes the worker use more
    than max_memory bytes, or take longer than timeout seconds, stops the worker instead of
    exhausting the machine (which takes every result of the run with it); the document counts
    as failed and a new worker takes the next one. Times are measured inside the worker."""

    def __init__(self, name: str, product: str | None = None, use_vlm: bool = True,
                 max_memory: int | None = None, timeout: float = 1800, startup_timeout: float = 3600) -> None:
        self.name = name
        self.product = product
        self.use_vlm = use_vlm
        self.max_memory = max_memory or int(_memory_total() * 0.75) or 2**62
        self.timeout = timeout
        self.startup_timeout = startup_timeout
        self._proc = None
        self._conn = None

    def _start(self) -> None:
        import multiprocessing

        ctx = multiprocessing.get_context("spawn")
        self._conn, child = ctx.Pipe()
        self._proc = ctx.Process(target=_worker, args=(self.name, self.product, self.use_vlm, child), daemon=True)
        self._proc.start()
        child.close()
        self._wait(self.startup_timeout, "starting")

    def _stop(self) -> None:
        if self._proc is not None and self._proc.is_alive():
            self._proc.kill()
            self._proc.join(30)
        self._proc = None

    def _wait(self, timeout: float, what: str):
        deadline = time.monotonic() + timeout
        while True:
            if self._conn.poll(0.5):
                kind, payload = self._conn.recv()
                if kind == "error":
                    raise RuntimeError(payload)
                return payload
            if not self._proc.is_alive():
                code = self._proc.exitcode
                self._proc = None
                raise RuntimeError(f"{self.name} worker ended while {what} (exit code {code})")
            used = _rss(self._proc.pid)
            if used > self.max_memory:
                self._stop()
                raise MemoryError(f"{self.name} used {used / 2**30:.1f} GB while {what} "
                                  f"(limit {self.max_memory / 2**30:.1f} GB)")
            if time.monotonic() > deadline:
                self._stop()
                raise TimeoutError(f"{self.name} took more than {timeout:.0f} s while {what}")

    def warm_up(self) -> None:
        if self._proc is None:
            self._start()

    def convert(self, path: Path) -> Converted:
        self.warm_up()
        self._conn.send(str(path))
        markdown, stats = self._wait(self.timeout, f"converting {path.name}")
        return Converted(markdown, DocumentStats(**stats))

    def close(self) -> None:
        if self._proc is not None and self._proc.is_alive():
            try:
                self._conn.send(None)
                self._proc.join(30)
            except OSError:
                pass
        self._stop()


def make_system(name: str, settings: Settings, use_vlm: bool = True, pipeline=None):
    product = settings.product_info
    if name == "docextract":
        return DocextractSystem(settings, pipeline, use_vlm)
    if name == "ppstructure":
        return PPStructureSystem(product.paddle_lang)
    if name == "tesseract":
        return TesseractSystem(product.tesseract_langs)
    if name == "text_layer":
        return TextLayerSystem()
    raise ValueError(f"unknown system {name!r}: one of {', '.join(SYSTEMS)}")


__all__ = ["Converted", "IsolatedSystem", "SYSTEMS", "make_system"]
