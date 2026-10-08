import io
import json

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from docextract import DocumentPipeline, ExtractOptions
from docextract.engines.base import Engines
from docextract.models import Method, PageKind, RegionType, ValidationStatus

from .conftest import FakeLayout, FakeOCR, FakeTable, FakeVLM, make_digital_pdf, make_scanned_pdf


def test_digital_pdf_without_models(settings):
    """No PaddleOCR, no VLM: the PDF's own structure is enough for a digital PDF."""
    result = DocumentPipeline(settings, Engines()).process_bytes(make_digital_pdf(), "report.pdf")
    md = result.markdown
    assert "# ACME Annual Report" in md
    assert "## 1. Introduction" in md and "## 2. Outlook" in md
    assert "| Item | Q1 | Q2 |" in md and "| 10.5 | 11.5 | 12.5 |" in md
    assert "1,234.3 million USD, up 12.3%" in md
    assert result.pages[0].kind == PageKind.DIGITAL
    methods = {r.method for r in result.iter_regions()}
    assert methods <= {Method.PDF_TEXT, Method.PDF_TABLE}
    table = next(r for r in result.iter_regions() if r.type == RegionType.TABLE)
    assert table.status == ValidationStatus.PASSED and table.html.startswith("<table>")
    assert table.source.page == 1 and table.bbox.x0 == 72.0
    # page number is kept in JSON but left out of the Markdown
    assert any(r.type == RegionType.PAGE_NUMBER for r in result.iter_regions())
    data = json.loads(result.to_json())
    assert data["source"]["format"] == "pdf" and data["stats"]["regions"] == len(list(result.iter_regions()))


def test_scanned_pdf_routing_and_fallback(settings, tmp_path):
    vlm = FakeVLM()
    ocr = FakeOCR()
    ragged = "<table><tr><td>Chỉ tiêu</td><td>Năm 2024</td></tr><tr><td>Doanh thu</td></tr></table>"
    engines = Engines(layout=FakeLayout(), ocr=ocr, table=FakeTable(html=ragged), vlm=vlm)
    settings = settings.model_copy(update={"cache_dir": str(tmp_path / "cache")})
    pipeline = DocumentPipeline(settings, engines)
    result = pipeline.process_bytes(make_scanned_pdf(), "scan.pdf")

    regions = {r.type: r for r in result.iter_regions()}
    assert result.pages[0].kind == PageKind.SCANNED
    assert regions[RegionType.TITLE].method == Method.OCR
    assert regions[RegionType.TEXT].content.startswith("Doanh thu năm 2025")
    # table: recogniser output is ragged -> escalated to the VLM, cross-checked against OCR numbers
    table = regions[RegionType.TABLE]
    assert [a.method for a in table.attempts] == [Method.TABLE_RECOGNITION, Method.VLM]
    assert table.method == Method.VLM and table.status == ValidationStatus.PASSED
    assert regions[RegionType.IMAGE].method == Method.VLM
    # the VLM never transcribes text
    assert {r.task for r in vlm.requests} <= {"table", "image", "chart", "formula"}
    assert result.stats.escalations == 1 and result.stats.vlm_calls == 2

    # second run is served from the cache
    calls = len(vlm.requests)
    again = pipeline.process_bytes(make_scanned_pdf(), "scan.pdf")
    assert len(vlm.requests) == calls
    assert again.stats.cache_hits > 0
    assert again.markdown == result.markdown


def test_ocr_dropout_is_flagged_not_sent_to_vlm(settings):
    texts = {
        60: ("BÁO CÁO", 0.99),
        110: ("Cng hòa xã hi ch nghĩa Vit Nam Đc lp T do Hnh phúc", 0.98),
        160: ("", 0.0),
        210: ("", 0.0),
    }
    vlm = FakeVLM()
    engines = Engines(layout=FakeLayout(), ocr=FakeOCR(texts), vlm=vlm)
    result = DocumentPipeline(settings, engines).process_bytes(make_scanned_pdf(), "scan.pdf")
    text = next(r for r in result.iter_regions() if r.type == RegionType.TEXT)
    assert text.method == Method.OCR and text.status == ValidationStatus.NEEDS_REVIEW
    assert "cần kiểm tra" in result.markdown
    assert all(r.task != "text" for r in vlm.requests)


def test_scanned_page_without_layout_model_is_ocred_whole(settings):
    engines = Engines(ocr=FakeOCR({250: ("Trang trắng", 0.99), 260: ("Trang trắng", 0.99)}))
    image = Image.new("RGB", (400, 300), "white")
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    result = DocumentPipeline(settings, engines).process_bytes(buf.getvalue(), "page.png")
    (region,) = list(result.iter_regions())
    assert region.meta.get("full_page") and region.method == Method.OCR
    assert result.source.format == "image" and result.pages[0].unit == "px"


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Báo cáo tổng kết", level=0)
    doc.add_heading("1. Kết quả", level=1)
    doc.add_paragraph("Doanh thu đạt 1.234,5 tỷ đồng.")
    doc.add_paragraph("Mục thứ nhất", style="List Bullet")
    doc.add_paragraph("Mục thứ hai", style="List Bullet")
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Gộp"
    table.cell(0, 2).text = "C"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Dọc"
    for r in (1, 2):
        table.cell(r, 1).text = f"{r}.1"
        table.cell(r, 2).text = f"{r}.2"
    buf = io.BytesIO()
    image = Image.new("RGB", (200, 120), (200, 30, 30))
    img_buf = io.BytesIO()
    image.save(img_buf, format="PNG")
    doc.add_picture(img_buf)
    doc.save(buf)
    return buf.getvalue()


def test_docx_structure(settings):
    vlm = FakeVLM()
    result = DocumentPipeline(settings, Engines(vlm=vlm)).process_bytes(_docx_bytes(), "baocao.docx")
    md = result.markdown
    assert "# Báo cáo tổng kết" in md
    assert "## 1. Kết quả" in md
    assert "- Mục thứ nhất\n- Mục thứ hai" in md
    assert 'colspan="2"' in md and 'rowspan="2"' in md
    assert "**[Hình ảnh]** Biểu tượng" in md
    assert result.source.format == "docx" and result.pages[0].number is None
    table = next(r for r in result.iter_regions() if r.type == RegionType.TABLE)
    assert table.method == Method.DOCX and table.source.locator == "body/tbl[1]"


def test_docx_image_without_vlm_keeps_alt_text(settings):
    result = DocumentPipeline(settings, Engines()).process_bytes(_docx_bytes(), "a.docx", ExtractOptions(use_vlm=False))
    image = next(r for r in result.iter_regions() if r.type == RegionType.IMAGE)
    assert image.method == Method.NONE and image.status == ValidationStatus.NEEDS_REVIEW


def test_page_selection(settings):
    result = DocumentPipeline(settings, Engines()).process_bytes(
        make_digital_pdf(), "r.pdf", ExtractOptions(pages=[1, 5])
    )
    assert [p.number for p in result.pages] == [1]
